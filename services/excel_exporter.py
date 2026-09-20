import pandas as pd
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter

def format_sheet(ws):
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")
    header_font = Font(bold=True)
    
    # Format header
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    
    # Auto-adjust column width and number formats
    for col in ws.columns:
        max_length = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if isinstance(cell.value, (int, float)):
                cell.number_format = '#,##0'
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
            # Wrap text cho diễn giải và ghi chú
            if cell.row > 1 and ws.cell(row=1, column=cell.column).value in ['Diễn giải', 'Diễn giải ngân hàng', 'Diễn giải kế toán', 'Ghi chú']:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                
        adjusted_width = min(max_length + 2, 50)
        ws.column_dimensions[col_letter].width = adjusted_width

def export_results(matched_groups, unmatched_bank, unmatched_acc, output_path):
    # --- SHEET KHOP ---
    khop_rows = []
    for grp in matched_groups:
        b_len = len(grp['bank'])
        a_len = len(grp['acc'])
        
        bank_total = sum(b['amount'] for b in grp['bank'])
        bank_desc_concat = " | ".join([b['description'] for b in grp['bank']])
        note = f"{grp['type']}. {b_len} giao dịch ngân hàng đối ứng với {a_len} dòng kế toán."
        
        # Ghi đầy đủ các dòng kế toán vào sheet KHOP
        for a in grp['acc']:
            khop_rows.append({
                'Ngày': a['date'],
                'Diễn giải': a['description'],  # Diễn giải gốc
                'Nợ': a['debit'],
                'Có': a['credit'],
                'Đối ứng ngân hàng': bank_total,
                'Trạng thái': grp['type'],
                'Ghi chú': note
            })
            
    df_khop = pd.DataFrame(khop_rows)
    
    # --- SHEET KHONG_KHOP ---
    khong_khop_rows = []
    
    # Xử lý ngân hàng thừa (không có kế toán đối ứng)
    for b in unmatched_bank:
        khong_khop_rows.append({
            'Ngày ngân hàng': b['date'],
            'Diễn giải ngân hàng': b['description'],
            'Tiền ngân hàng': b['amount'],
            'Ngày kế toán': '',
            'Diễn giải kế toán': '',
            'Tiền kế toán': 0,
            'Chênh lệch': b['amount'],
            'Trạng thái': 'Chưa nhập / Ngân hàng thừa',
            'Ghi chú': 'Đã dò nhưng không tìm thấy dòng kế toán phù hợp. Khả năng giao dịch này chưa được hạch toán.'
        })
        
    # Xử lý kế toán thừa (không có ngân hàng đối ứng)
    for a in unmatched_acc:
        khong_khop_rows.append({
            'Ngày ngân hàng': '',
            'Diễn giải ngân hàng': '',
            'Tiền ngân hàng': 0,
            'Ngày kế toán': a['date'],
            'Diễn giải kế toán': a['description'],
            'Tiền kế toán': a['amount'],
            'Chênh lệch': -a['amount'],
            'Trạng thái': 'Thừa / Kế toán đã ghi',
            'Ghi chú': 'Đã dò nhưng không tìm thấy giao dịch ngân hàng tương ứng. Cần kiểm tra lại chứng từ hoặc hạch toán nhầm.'
        })
        
    df_khong_khop = pd.DataFrame(khong_khop_rows)
    
    # Ghi Excel
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        if not df_khop.empty:
            df_khop.to_excel(writer, sheet_name='KHOP', index=False)
        else:
            pd.DataFrame(columns=['Ngày', 'Diễn giải', 'Nợ', 'Có', 'Đối ứng ngân hàng', 'Trạng thái', 'Ghi chú']).to_excel(writer, sheet_name='KHOP', index=False)
            
        if not df_khong_khop.empty:
            df_khong_khop.to_excel(writer, sheet_name='KHONG_KHOP', index=False)
        else:
            pd.DataFrame(columns=['Ngày ngân hàng', 'Diễn giải ngân hàng', 'Tiền ngân hàng', 'Ngày kế toán', 'Diễn giải kế toán', 'Tiền kế toán', 'Chênh lệch', 'Trạng thái', 'Ghi chú']).to_excel(writer, sheet_name='KHONG_KHOP', index=False)
            
        format_sheet(writer.sheets['KHOP'])
        format_sheet(writer.sheets['KHONG_KHOP'])